#!/usr/bin/env python3
"""
ficheiro: generate_docs.py

Consolida os arquivos Markdown da pasta docs/ em um único documento Word
(docs/Revisao_Documentacao.docx), pronto para revisão/edição no Microsoft
Word antes de exportar para PDF.

Suporta o subconjunto de Markdown usado nos documentos deste projeto:
títulos (#, ##, ###), parágrafos com **negrito**, `código inline` e
[links](url), listas (- / * / 1.), blocos de código cercados por ```
(com nota especial para blocos ```mermaid), tabelas estilo GFM (|...|),
citações (> texto) e separadores horizontais (---).

Não depende de nenhuma biblioteca de parsing de Markdown de terceiros —
apenas de `python-docx` — para manter a instalação mínima:

    pip install python-docx
    python generate_docs.py

Saída: docs/Revisao_Documentacao.docx
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor
except ImportError:
    sys.exit(
        "Dependência ausente: instale com  pip install python-docx"
    )

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "docs"
OUTPUT_PATH = DOCS_DIR / "Revisao_Documentacao.docx"

# (arquivo em docs/, título do capítulo no documento consolidado)
CHAPTER_ORDER: list[tuple[str, str]] = [
    ("modelo_dados.md", "1. Modelo de Dados"),
    ("modelo_c4.md", "2. Arquitetura — Modelo C4"),
    ("contrato_api.md", "3. Contrato da API REST"),
    ("ADRs.md", "4. Registros de Decisão Arquitetural (ADRs)"),
    ("deploy.md", "5. Guia de Deploy"),
]

FONT_BODY = "Calibri"
FONT_MONO = "Consolas"
COLOR_TITLE = RGBColor(0x1F, 0x38, 0x64)
COLOR_HEADING = RGBColor(0x2E, 0x54, 0x96)
COLOR_MUTED = RGBColor(0x59, 0x59, 0x59)
FILL_CODE = "F2F2F2"
FILL_HEADER_ROW = "2E5496"
FILL_QUOTE = "FBF3E7"
BORDER_COLOR = "BFBFBF"

INLINE_PATTERN = re.compile(
    r"(?P<link>\[(?P<link_text>[^\]]+)\]\((?P<link_url>[^)]+)\))"
    r"|(?P<bold>\*\*(?P<bold_text>.+?)\*\*)"
    r"|(?P<code>`(?P<code_text>[^`]+)`)"
    r"|(?P<italic>\*(?P<italic_text>[^*]+)\*)"
)

FILE_LABEL_RE = re.compile(r"^<!--\s*ficheiro:.*-->$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?$")


# ---------------------------------------------------------------------------
# Modelo de blocos intermediário
# ---------------------------------------------------------------------------


@dataclass
class Block:
    kind: str  # heading | paragraph | bullet_list | ordered_list | code | table | quote | hr
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    lang: str = ""
    code_lines: list[str] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


def parse_markdown(text: str) -> list[Block]:
    """Converte o texto Markdown em uma lista de blocos estruturados."""
    lines = text.splitlines()
    blocks: list[Block] = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Linha em branco ou label de arquivo: ignora
        if not stripped or FILE_LABEL_RE.match(stripped):
            i += 1
            continue

        # Heading
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            blocks.append(
                Block(
                    kind="heading",
                    level=len(heading_match.group(1)),
                    text=heading_match.group(2).strip(),
                )
            )
            i += 1
            continue

        # Horizontal rule
        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            blocks.append(Block(kind="hr"))
            i += 1
            continue

        # Fenced code block
        fence_match = re.match(r"^```\s*(\S*)$", stripped)
        if fence_match:
            lang = fence_match.group(1)
            code_lines: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # pula a cerca de fechamento
            blocks.append(Block(kind="code", lang=lang, code_lines=code_lines))
            continue

        # Blockquote (agrupa linhas consecutivas)
        if stripped.startswith(">"):
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            blocks.append(Block(kind="quote", text=" ".join(quote_lines)))
            continue

        # Table (linha com '|' seguida de separador com '-')
        if "|" in stripped and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1].strip()):
            header_cells = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2  # pula cabeçalho + separador
            rows: list[list[str]] = []
            while i < n and "|" in lines[i].strip() and lines[i].strip():
                row_cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row_cells)
                i += 1
            blocks.append(Block(kind="table", headers=header_cells, rows=rows))
            continue

        # Unordered list
        if re.match(r"^[-*]\s+", stripped):
            items = []
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(Block(kind="bullet_list", items=items))
            continue

        # Ordered list
        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append(Block(kind="ordered_list", items=items))
            continue

        # Paragraph: acumula linhas até uma linha em branco ou início de outro bloco
        para_lines = [stripped]
        i += 1
        while i < n and lines[i].strip() and not _starts_new_block(lines[i]):
            para_lines.append(lines[i].strip())
            i += 1
        blocks.append(Block(kind="paragraph", text=" ".join(para_lines)))

    return blocks


def _starts_new_block(line: str) -> bool:
    s = line.strip()
    return bool(
        re.match(r"^#{1,6}\s+", s)
        or re.match(r"^```", s)
        or re.match(r"^[-*]\s+", s)
        or re.match(r"^\d+\.\s+", s)
        or s.startswith(">")
        or re.match(r"^-{3,}$", s)
        or "|" in s
    )


# ---------------------------------------------------------------------------
# Helpers de baixo nível do python-docx
# ---------------------------------------------------------------------------


def _set_cell_background(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_borders(cell, color: str = BORDER_COLOR) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def _set_paragraph_shading(paragraph, hex_color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def _add_inline_runs(
    paragraph,
    text: str,
    base_size: int = 11,
    force_italic: bool = False,
    base_color: RGBColor | None = None,
) -> None:
    """Adiciona `text` ao parágrafo, interpretando **negrito**, *itálico*,
    `código` e [links](url). Quando `force_italic`/`base_color` são
    passados (uso em citações), aplica-os a todos os runs gerados, além
    da formatação própria de cada trecho (negrito continua em negrito
    dentro de uma citação, por exemplo)."""

    def _style(run, mono: bool = False) -> None:
        run.font.name = FONT_MONO if mono else FONT_BODY
        run.font.size = Pt(base_size - 1 if mono else base_size)
        if force_italic:
            run.italic = True
        if base_color is not None:
            run.font.color.rgb = base_color

    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            _style(paragraph.add_run(text[pos : match.start()]))

        if match.group("link"):
            run = paragraph.add_run(match.group("link_text"))
            _style(run)
            run.font.underline = True
            if base_color is None:
                run.font.color.rgb = RGBColor(0x2E, 0x54, 0x96)
        elif match.group("bold"):
            run = paragraph.add_run(match.group("bold_text"))
            _style(run)
            run.bold = True
        elif match.group("code"):
            run = paragraph.add_run(match.group("code_text"))
            _style(run, mono=True)
        elif match.group("italic"):
            run = paragraph.add_run(match.group("italic_text"))
            _style(run)
            run.italic = True

        pos = match.end()

    if pos < len(text):
        _style(paragraph.add_run(text[pos:]))


# ---------------------------------------------------------------------------
# Renderização dos blocos no Document
# ---------------------------------------------------------------------------

HEADING_STYLE_BY_LEVEL = {1: "Heading 2", 2: "Heading 3", 3: "Heading 4", 4: "Heading 5"}


def render_blocks(doc: Document, blocks: Iterable[Block]) -> None:
    for block in blocks:
        if block.kind == "heading":
            style = HEADING_STYLE_BY_LEVEL.get(block.level, "Heading 5")
            p = doc.add_paragraph(style=style)
            _add_inline_runs(p, block.text, base_size=13)

        elif block.kind == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            _add_inline_runs(p, block.text)

        elif block.kind == "bullet_list":
            for item in block.items:
                p = doc.add_paragraph(style="List Bullet")
                _add_inline_runs(p, item)

        elif block.kind == "ordered_list":
            for item in block.items:
                p = doc.add_paragraph(style="List Number")
                _add_inline_runs(p, item)

        elif block.kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            _set_paragraph_shading(p, FILL_QUOTE)
            _add_inline_runs(
                p, block.text, base_size=10.5, force_italic=True, base_color=COLOR_MUTED
            )

        elif block.kind == "hr":
            p = doc.add_paragraph()
            p_fmt = p.paragraph_format
            p_fmt.space_before = Pt(4)
            p_fmt.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), BORDER_COLOR)
            border.append(bottom)
            pPr.append(border)

        elif block.kind == "code":
            if block.lang == "mermaid":
                note = doc.add_paragraph()
                run = note.add_run(
                    "\u25b8 Diagrama Mermaid — visualize em mermaid.live, "
                    "num Markdown viewer compatível (ex.: GitHub, VS Code) "
                    "ou renderize antes de exportar o PDF final."
                )
                run.italic = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_MUTED

            code_text = "\n".join(block.code_lines) if block.code_lines else " "
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.left_indent = Cm(0.3)
            _set_paragraph_shading(p, FILL_CODE)
            for idx, line in enumerate(code_text.split("\n")):
                if idx > 0:
                    p.add_run().add_break()
                run = p.add_run(line if line else " ")
                run.font.name = FONT_MONO
                run.font.size = Pt(9.5)

        elif block.kind == "table":
            n_cols = len(block.headers)
            table = doc.add_table(rows=1, cols=n_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            header_cells = table.rows[0].cells
            for col_idx, header_text in enumerate(block.headers):
                cell = header_cells[col_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(header_text)
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)
                run.font.name = FONT_BODY
                _set_cell_background(cell, FILL_HEADER_ROW)
                _set_cell_borders(cell)

            for row_data in block.rows:
                row_cells = table.add_row().cells
                for col_idx in range(n_cols):
                    text = row_data[col_idx] if col_idx < len(row_data) else ""
                    cell = row_cells[col_idx]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    _add_inline_runs(p, text, base_size=10)
                    _set_cell_borders(cell)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ---------------------------------------------------------------------------
# Documento principal
# ---------------------------------------------------------------------------


def build_document() -> Document:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(11)

    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # Capa
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AtendSaúde")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = COLOR_TITLE
    run.font.name = FONT_BODY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Revisão de Documentação Técnica")
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_HEADING

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(
        "Desafio Técnico Full Stack — Seleção V-Lab (CIn-UFPE)\n"
        "Documento gerado automaticamente por generate_docs.py a partir de /docs — "
        "revise e ajuste livremente antes de exportar em PDF."
    )
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

    missing: list[str] = []
    for filename, chapter_title in CHAPTER_ORDER:
        path = DOCS_DIR / filename
        if not path.exists():
            missing.append(filename)
            continue

        heading = doc.add_paragraph(style="Heading 1")
        run = heading.add_run(chapter_title)
        run.font.color.rgb = COLOR_TITLE

        blocks = parse_markdown(path.read_text(encoding="utf-8"))
        render_blocks(doc, blocks)
        doc.add_page_break()

    if missing:
        print(f"[aviso] arquivos não encontrados em {DOCS_DIR}: {', '.join(missing)}")

    return doc


def main() -> None:
    if not DOCS_DIR.exists():
        sys.exit(f"Pasta não encontrada: {DOCS_DIR}")

    document = build_document()
    document.save(OUTPUT_PATH)
    print(f"Gerado: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
